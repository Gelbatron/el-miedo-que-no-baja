"""
Respaldo sin internet de la presentación interactiva (/presentar del visor).

Captura cada lámina del modo pitch en su último paso y arma:
- Pitch_El_miedo_que_no_baja_The_data_squad.pdf   (para proyectar desde una USB);
- Pitch_El_miedo_que_no_baja_The_data_squad.pptx  (mismas láminas, con las notas del orador);
- guion_pitch.md                                   (quién dice qué; conserva las preguntas del jurado).

El guion sale de la misma fuente que las notas del visor (página /presentar/guion), así que no se separan.

Uso (Windows, con Microsoft Edge):
    cd visor && npm run build && set VISOR_SIN_CLAVE=1 && npx next start -p 3210
    python docs/pitch/capturar_presentacion.py
"""
import html
import json
import re
import subprocess
import tempfile
import urllib.request
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

RAIZ = Path(__file__).resolve().parents[2]
OUT = RAIZ / "docs" / "pitch"
BASE = "http://localhost:3210"
EDGE = Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe")
NOMBRE = "Pitch_El_miedo_que_no_baja_The_data_squad"


def guion() -> list[dict]:
    pagina = urllib.request.urlopen(f"{BASE}/presentar/guion", timeout=30).read().decode("utf-8")
    m = re.search(r'<script type="application/json" id="guion-datos">(.*?)</script>', pagina, re.S)
    if not m:
        raise SystemExit("No encontré el guion en /presentar/guion. ¿Está corriendo el visor con VISOR_SIN_CLAVE=1?")
    return json.loads(html.unescape(m.group(1)))


def capturar(diapo: dict, carpeta: Path) -> Path:
    ruta = carpeta / f"{diapo['id']}.png"
    url = f"{BASE}/presentar#{diapo['id']}~{diapo['pasos'] - 1}~c"
    subprocess.run([str(EDGE), "--headless=new", "--disable-gpu", "--hide-scrollbars", f"--user-data-dir={carpeta / 'perfil'}",
                    "--window-size=1920,1080", "--virtual-time-budget=9000", f"--screenshot={ruta}", url],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    return ruta


def main():
    diapos = [d for d in guion() if not d["soloClase"]]
    with tempfile.TemporaryDirectory() as tmp:
        carpeta = Path(tmp)
        imagenes = [Image.open(capturar(d, carpeta)).convert("RGB") for d in diapos]
        pdf = OUT / f"{NOMBRE}.pdf"
        imagenes[0].save(pdf, save_all=True, append_images=imagenes[1:], resolution=144)

        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for d, img in zip(diapos, imagenes):
            png = carpeta / f"{d['id']}_pptx.png"
            img.save(png)
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_picture(str(png), 0, 0, prs.slide_width, prs.slide_height)
            s.notes_slide.notes_text_frame.text = f"{d['quien']} ({d['tiempo']}): {d['nota']}"
        pptx = OUT / f"{NOMBRE}.pptx"
        prs.save(pptx)

    md = OUT / "guion_pitch.md"
    preguntas = ""
    if md.exists():
        actual = md.read_text(encoding="utf-8")
        if "## Preguntas" in actual:
            preguntas = actual[actual.index("## Preguntas"):]
    filas = "\n".join(f"| {k} · {d['titulo']} | {d['tiempo']} | {d['quien']} | {d['nota']} |" for k, d in enumerate(diapos, 1))
    md.write_text(
        "# Guion del pitch · 6 minutos\n\n"
        "**El miedo que no baja** · The data squad. Se presenta desde el visor (botón ▶ Presentar, modo «Pitch 6 min»). "
        "→ avanza un paso, N muestra estas notas, D abre «¿Cómo lo sabemos?» para las preguntas. "
        f"Respaldo sin internet: `{NOMBRE}.pdf` en una USB. Ensayar con el cronómetro hasta cerrar en 5:30.\n\n"
        "| Lámina | Tiempo | Quién | Qué decir (idea, no para leer) |\n|---|---|---|---|\n"
        f"{filas}\n\n"
        "En la clase completa se suman «¿Qué predice mejor el miedo?» (validación cruzada) y el simulador.\n\n"
        f"{preguntas}",
        encoding="utf-8")
    print(f"{len(diapos)} láminas: {pdf.name}, {pptx.name} y {md.name}")


if __name__ == "__main__":
    main()
