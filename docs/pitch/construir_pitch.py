"""Presentación del pitch (6 min) con la identidad del visor: franja roja y oro, títulos en rojo, oro = lo que se siente,
azul = lo que se registra. Las cifras salen de outputs/resultados.json y las notas del orador del guion.

Uso:  python docs/pitch/construir_pitch.py   (después de correr src/pipeline.py)
Tipografía: Segoe UI / Segoe UI Black, presentes en Windows, para que el .pptx y el .pdf se vean igual en cualquier equipo.
"""
import json
from pathlib import Path

import qrcode
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

R = Path(__file__).resolve().parents[2]
S = json.loads((R / "outputs" / "resultados.json").read_text(encoding="utf-8"))
FIG = R / "outputs" / "figuras"
OUT = R / "docs" / "pitch"
URL = "https://el-miedo-que-no-baja.vercel.app"

hexa = lambda h: RGBColor.from_string(h.lstrip("#").upper())
TINTA, TINTA2, TINTA3 = hexa("#1e1e1e"), hexa("#4b4b4b"), hexa("#767676")
ROJO, VINO, ORO, ORO_CLARO = hexa("#bf0a00"), hexa("#ab0f24"), hexa("#e9a924"), hexa("#ffd98a")
PERC, PERC_T, REG, REG_T, MUJ = hexa("#d99a16"), hexa("#8a5a00"), hexa("#3366cc"), hexa("#1f4a9e"), hexa("#8e44c8")
FONDO, BLANCO = hexa("#f2f2f2"), hexa("#ffffff")
CUERPO, TITULAR = "Segoe UI", "Segoe UI Black"
f = lambda x, d=1: f"{x:,.{d}f}".replace(",", "§").replace(".", ",").replace("§", ".").replace("-", "−")

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIO = prs.slide_layouts[6]
W = 13.333


def texto(slide, x, y, w, h, contenido, tam=18, color=TINTA, negrita=False, alinear=PP_ALIGN.LEFT, ancla=MSO_ANCHOR.TOP,
          interlineado=1.1, fuente=CUERPO):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ancla
    tf.margin_left = tf.margin_right = Inches(0.02)
    for i, linea in enumerate(contenido if isinstance(contenido, list) else [contenido]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = alinear
        p.line_spacing = interlineado
        est = {}
        for frag, est in (linea if isinstance(linea, list) else [(linea, {})]):
            r = p.add_run()
            r.text = frag
            r.font.name = est.get("fuente", fuente)
            r.font.size = Pt(est.get("tam", tam))
            r.font.bold = est.get("negrita", negrita)
            r.font.color.rgb = est.get("color", color)
        p.space_after = Pt(est.get("despues", 6))
    return tb


def rect(slide, x, y, w, h, color, redondeo=0.06, forma=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
    if forma == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = redondeo
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def franja(slide):
    """Franja superior del visor: 72 % rojo, 28 % oro."""
    rect(slide, 0, 0, W * 0.72, 0.09, ROJO, forma=MSO_SHAPE.RECTANGLE)
    rect(slide, W * 0.72, 0, W * 0.28, 0.09, ORO, forma=MSO_SHAPE.RECTANGLE)


def tarjeta_cifra(slide, x, y, w, h, etiqueta, cifra, detalle, color_borde, color_cifra):
    rect(slide, x, y, w, h, FONDO)
    rect(slide, x, y, w, 0.09, color_borde, forma=MSO_SHAPE.RECTANGLE)
    texto(slide, x + 0.25, y + 0.22, w - 0.5, 0.35, etiqueta.upper(), tam=10.5, negrita=True, color=color_cifra)
    texto(slide, x + 0.25, y + 0.5, w - 0.5, 1.0, cifra, tam=50, color=color_cifra, fuente=TITULAR)
    texto(slide, x + 0.25, y + h - 0.85, w - 0.5, 0.8, detalle, tam=12.5, color=TINTA2)


def notas(slide, quien, texto_notas):
    slide.notes_slide.notes_text_frame.text = f"{quien}: {texto_notas}"


def base(antetitulo, titulo, n):
    s = prs.slides.add_slide(VACIO)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BLANCO
    franja(s)
    texto(s, 0.6, 0.38, 9, 0.35, antetitulo.upper(), tam=11, negrita=True, color=ROJO)
    texto(s, 0.6, 0.68, 12.1, 0.9, titulo, tam=27, color=ROJO, fuente=TITULAR)
    texto(s, 11.6, 7.0, 1.2, 0.3, f"{n} / 8", tam=10, color=TINTA3, alinear=PP_ALIGN.RIGHT)
    texto(s, 0.6, 7.0, 9, 0.3, "El miedo que no baja · The data squad · Bogotá DataJam 5", tam=10, color=TINTA3)
    return s


def imagen(s, ruta, x, y, w=None, h=None):
    return s.shapes.add_picture(str(ruta), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


a, g, v, auc = S["ame_hechos"], S["emv_mujer_ame"], S["ibs_vs_cifra_oculta"], S["auc"]

# 1. Portada (roja, como el encabezado del visor)
s = prs.slides.add_slide(VACIO)
s.background.fill.solid()
s.background.fill.fore_color.rgb = ROJO
rect(s, W * 0.72, 0, W * 0.28, 0.09, ORO, forma=MSO_SHAPE.RECTANGLE)
texto(s, 0.8, 0.8, 11.5, 0.4, "BOGOTÁ DATAJAM EDICIÓN 5 · SEGURIDAD CIUDADANA CON ENFOQUE DE GÉNERO", tam=12, negrita=True, color=ORO_CLARO)
texto(s, 0.8, 1.35, 11.5, 2.4, [[("El miedo", {"color": BLANCO})], [("que no baja", {"color": ORO})]], tam=68, fuente=TITULAR, interlineado=0.9)
rect(s, 0.85, 3.85, 1.4, 0.08, ORO, forma=MSO_SHAPE.RECTANGLE)
texto(s, 0.8, 4.15, 10.8, 1.4, "¿Por qué la inseguridad percibida no baja cuando el delito registrado sí, y en qué territorios debe intervenir primero el Distrito?",
      tam=22, color=BLANCO)
texto(s, 0.8, 5.85, 11.5, 1.0, [[("Equipo #16 · The data squad", {"negrita": True})],
                                "Rafael Estiven Londoño Camacho · Brandow Brusly León Rodríguez · Samuel Esteban Riveros Cubides"],
      tam=14, color=BLANCO)
notas(s, "Brandow", "Somos The data squad. Nuestra pregunta es por qué la gente en Bogotá no se siente más segura aunque las cifras "
                    "de delito bajan, y dónde debería actuar primero el Distrito. (0:00-0:20)")

# 2. El problema
s = base("El problema", "Bogotá registra menos delitos, pero no se siente más segura", 2)
tarjeta_cifra(s, 0.6, 1.8, 3.9, 2.35, "Lo que se cuenta · SCJ", f"{f(S['hurto_personas']['cambio_pct'], 0)} %",
              "hurto a personas registrado, enero-agosto 2019 → 2026", REG, REG_T)
tarjeta_cifra(s, 0.6, 4.35, 3.9, 2.35, "Lo que se siente · Bienal 2025", f"{f(S['ciudad_insegura_pct'], 0)} %",
              "de los adultos califica a Bogotá como insegura", PERC, PERC_T)
imagen(s, FIG / "fig1_serie_delitos.png", 4.9, 1.8, w=7.9)
notas(s, "Brandow", "Entre enero y agosto, el hurto a personas registrado cayó 65 % desde 2019. Pero el 74 % de los adultos dice que la "
                    "ciudad es insegura. La política se planea con el registro y la gente vive otra cosa. Señalar que la violencia "
                    "intrafamiliar y los delitos sexuales van en sentido contrario. (0:20-1:00)")

# 3. Datos y método
s = base("Datos y método", "Siete fuentes, una llave territorial y una escalera analítica", 3)
fuentes = [("SCJ", "Delitos de alto impacto 2018-2026 y llamadas a la Línea 123", REG),
           ("SDP", "Proyecciones de población 2005-2035 (tasas por 100.000)", REG),
           ("SCRD", f"Bienal de Cultura Ciudadana 2025 · {f(S['n_bienal'], 0)} adultos", PERC),
           ("EMV", f"Encuesta Mundial de Valores 2024 · {f(S['n_emv_bogota'], 0)} en Bogotá", PERC),
           ("UAESP · DADEP · IDECA", "Alumbrado, espacio público y límites", REG)]
for i, (k, d, c) in enumerate(fuentes):
    y = 1.85 + i * 0.95
    rect(s, 0.6, y, 5.6, 0.8, FONDO)
    rect(s, 0.6, y, 0.09, 0.8, c, forma=MSO_SHAPE.RECTANGLE)
    texto(s, 0.9, y + 0.08, 5.1, 0.3, k, tam=13, negrita=True)
    texto(s, 0.9, y + 0.4, 5.1, 0.4, d, tam=12, color=TINTA2)
pasos = [("Descriptivo", "Series ene-ago, tasas por 100.000 y porcentajes ponderados con error estándar"),
         ("Diagnóstico", "Logit ponderado, efectos marginales, IC por bootstrap de localidades, efectos fijos y wild bootstrap"),
         ("Predictivo", "AUC fuera de muestra por bloques de variables"),
         ("Prescriptivo", "Índice de Brecha con incertidumbre y simulación de escenarios")]
for i, (k, d) in enumerate(pasos):
    y = 1.85 + i * 1.18
    rect(s, 6.8, y, 0.6, 0.6, ORO, forma=MSO_SHAPE.OVAL)
    texto(s, 6.8, y + 0.07, 0.6, 0.5, str(i + 1), tam=18, color=TINTA, alinear=PP_ALIGN.CENTER, fuente=TITULAR)
    texto(s, 7.6, y + 0.0, 5.2, 0.35, k, tam=15, negrita=True)
    texto(s, 7.6, y + 0.4, 5.2, 0.75, d, tam=12.5, color=TINTA2)
notas(s, "Rafael", "Integramos siete fuentes públicas y las unimos por el código de localidad, con tasas por 100.000 habitantes. "
                   "Subimos la escalera completa: descripción, diagnóstico con un logit ponderado, predicción fuera de muestra y escenarios. (1:00-1:50)")

# 4. Hallazgo 1
s = base("Hallazgo 1 · cifra oculta", "Donde hay más miedo que delito, el delito no se registra", 4)
imagen(s, FIG / "fig5_brecha_vs_cifra_oculta.png", 0.5, 1.75, w=7.7)
rect(s, 8.55, 1.8, 4.25, 4.9, VINO)
texto(s, 8.85, 2.0, 3.8, 0.4, "CIFRA OCULTA", tam=11, negrita=True, color=ORO_CLARO)
texto(s, 8.85, 2.35, 3.8, 1.1, f"{f(S['cifra_oculta_bogota'])} a 1", tam=50, color=ORO, fuente=TITULAR)
texto(s, 8.85, 3.45, 3.75, 3.2, [
    "víctimas por cada delito de alto impacto registrado en Bogotá.",
    f"{S['cifra_oculta_max'][0]}: {f(S['cifra_oculta_max'][1])} · {S['cifra_oculta_min'][0]}: {f(S['cifra_oculta_min'][1])}.",
    f"Índice de Brecha frente a cifra oculta: Spearman {f(v['rho'], 2)} (IC 95 % {f(v['ic_inf'], 2)} a {f(v['ic_sup'], 2)}; p ajustado {f(v['p_bh'], 3)})."],
    tam=13.5, color=BLANCO)
notas(s, "Rafael", f"Por cada delito registrado hay unas {f(S['cifra_oculta_bogota'])} víctimas. En Bosa más de 10; en Los Mártires menos de 1. "
                   f"Donde hay más miedo que delito registrado, más delito queda sin registrar: correlación de {f(v['rho'], 2)}, la única que "
                   "sobrevive a la corrección por pruebas múltiples. El miedo de la periferia no es exagerado. (1:50-2:40)")

# 5. Hallazgo 2
s = base("Hallazgo 2 · entorno", "El barrio se siente por su entorno, no por el hurto registrado", 5)
imagen(s, FIG / "fig3_efectos_marginales.png", 0.5, 1.75, w=8.0)
tarjeta_cifra(s, 8.85, 1.8, 3.95, 2.3, "Atracos en el parque", f"+{f(a['parque_atracos']['ame_pp'])} pp",
              "más probabilidad de sentir insegura la zona", PERC, PERC_T)
rect(s, 8.85, 4.3, 3.95, 2.4, FONDO)
texto(s, 9.1, 4.45, 3.5, 2.2, [
    [("AUC fuera de muestra", {"negrita": True, "color": TINTA})],
    f"Solo perfil: {f(auc['Perfil (sexo, edad, estrato, tenencia)'], 3)}",
    [(f"+ delito registrado: {f(auc['Perfil + delito registrado'], 3)}", {"color": REG_T})],
    [(f"+ entorno: {f(auc['Perfil + hechos del entorno'], 3)} · todo: {f(auc['Todo'], 3)}", {"color": PERC_T})],
    "Hurto de la localidad: wild bootstrap p = 0,92."], tam=13, color=TINTA2)
notas(s, "Rafael", f"Ver atracos en el parque cercano suma {f(a['parque_atracos']['ame_pp'])} puntos; los conflictos vecinales, hasta "
                   f"{f(a['conflictos_vecinos']['ame_pp'])}. El hurto registrado no suma nada. Agregar el delito registrado no mejora la "
                   "predicción (0,609 a 0,611); el entorno la sube a 0,78. (2:40-3:30)")

# 6. Hallazgo 3
s = base("Hallazgo 3 · enfoque de género", "Las mujeres reportan menos delitos y recortan más su vida", 6)
imagen(s, FIG / "fig6_genero.png", 0.5, 2.1, w=7.8)
rect(s, 8.7, 1.8, 4.1, 4.3, FONDO)
rect(s, 8.7, 1.8, 4.1, 0.09, MUJ, forma=MSO_SHAPE.RECTANGLE)
texto(s, 8.95, 2.05, 3.7, 4.6, [
    [("A igual edad, estrato y victimización:", {"negrita": True, "color": TINTA})],
    [(f"+{f(g['teme_acoso_sexual']['ame_pp'])} pp", {"color": MUJ, "tam": 28, "fuente": TITULAR}), ("  temor al acoso sexual", {})],
    [(f"+{f(g['barrio_inseguro']['ame_pp'])} pp", {"color": MUJ, "tam": 28, "fuente": TITULAR}), ("  insegura en su barrio", {})],
    [(f"+{f(g['no_sale_noche']['ame_pp'])} pp", {"color": MUJ, "tam": 28, "fuente": TITULAR}), ("  evita salir de noche", {})],
    f"Mientras el hurto cae, la violencia intrafamiliar sube {f(S['vif_cambio_pct_2018_2026'])} % y los delitos sexuales "
    f"{f(S['delitos_sexuales_cambio_pct_2018_2026'])} % (2018-2026)."], tam=13.5, color=TINTA2)
notas(s, "Samuel", f"Las mujeres reportan menos delitos, pero a igual edad, estrato y victimización tienen {f(g['teme_acoso_sexual']['ame_pp'])} "
                   f"puntos más de temor al acoso sexual y {f(g['no_sale_noche']['ame_pp'])} más de evitar salir de noche. El costo del miedo "
                   "se paga en autonomía. (3:30-4:15)")

# 7. Dónde intervenir
s = base("Recomendación", "Dónde intervenir primero y cómo medirlo", 7)
imagen(s, FIG / "fig7_indice_brecha.png", 0.45, 1.75, w=6.0)
top5 = S["ibs_top5"]
recs = [("Medir con victimización", "Encuesta Distrital de Percepción por UPL junto al registro SCJ: corrige la asignación donde se subregistra."),
        ("Entornos que se sienten seguros", f"Parques sin atracos, mediación vecinal y alumbrado en {', '.join(top5[:4])} y {top5[4]}."),
        ("Autonomía de las mujeres", "«% que evita salir de noche, por sexo» como indicador de resultado."),
        ("Datos para evaluar", "Alumbrado georreferenciado (el del portal es de 2019) y delitos por UPZ.")]
for i, (k, d) in enumerate(recs):
    y = 1.75 + i * 1.24
    rect(s, 6.75, y, 6.05, 1.1, FONDO)
    rect(s, 6.92, y + 0.27, 0.55, 0.55, ORO, forma=MSO_SHAPE.OVAL)
    texto(s, 6.92, y + 0.32, 0.55, 0.5, str(i + 1), tam=16, color=TINTA, alinear=PP_ALIGN.CENTER, fuente=TITULAR)
    texto(s, 7.65, y + 0.1, 5.0, 0.35, k, tam=15, negrita=True, color=ROJO)
    texto(s, 7.65, y + 0.47, 5.0, 0.6, d, tam=12, color=TINTA2)
notas(s, "Brandow", f"Cuatro recomendaciones, una frase cada una: medir con victimización por UPL; entornos que se sienten seguros en "
                    f"{', '.join(top5)}; la autonomía de las mujeres como indicador; y publicar alumbrado georreferenciado y delitos por UPZ. (4:15-5:10)")

# 8. Indicador, visor y cierre
s = base("Indicador y visor", "Un indicador medible, comparable y con fuentes que ya existen", 8)
texto(s, 0.6, 1.8, 7.6, 3.9, [
    [("Fórmula: ", {"negrita": True, "color": TINTA}), ("z(% que siente insegura su zona) − z(ln delitos de alto impacto por 100.000)", {})],
    [("Fuentes: ", {"negrita": True, "color": TINTA}), ("EDP o Bienal (percepción), SCJ (delito), SDP (población)", {})],
    [("Periodicidad: ", {"negrita": True, "color": TINTA}), ("semestral por sector de UPL, anual por localidad y UPL", {})],
    [("Línea base 2025: ", {"negrita": True, "color": TINTA}), (f"{f(S['zona_insegura_pct'])} % siente insegura su zona", {})],
    [("Simulador: ", {"negrita": True, "color": TINTA}), ("con la mitad de las zonas y parques mejorados, −7,4 pp en Bogotá", {})],
    [("Todo reproducible: ", {"negrita": True, "color": TINTA}), ("notebook, repositorio, Power BI, Excel y descargas CSV y GeoJSON", {})]],
    tam=15, color=TINTA2, interlineado=1.12)
OUT.mkdir(parents=True, exist_ok=True)
qr_ruta = OUT / "qr_visor.png"
qrcode.make(URL).save(qr_ruta)
imagen(s, qr_ruta, 9.55, 1.75, w=2.7)
texto(s, 8.6, 4.45, 4.6, 0.8, [[(URL.removeprefix("https://"), {"negrita": True, "color": TINTA})], "clave en el correo de entrega"],
      tam=12.5, color=TINTA2, alinear=PP_ALIGN.CENTER)
rect(s, 0.6, 5.75, 12.15, 1.05, ROJO)
texto(s, 0.9, 5.75, 11.6, 1.05, [[("Bogotá ya produce los datos para cerrar esta brecha. ", {"color": BLANCO}), ("Hay que leerlos juntos.", {"color": ORO})]],
      tam=22, fuente=TITULAR, ancla=MSO_ANCHOR.MIDDLE)
notas(s, "Samuel", "Proponemos el Índice de Brecha de Seguridad: se calcula cada semestre con fuentes que ya existen. Todo está en el visor "
                   "(QR) y en el repositorio, reproducible de punta a punta. Si sobra tiempo, abrir el simulador. Cierre (Brandow): Bogotá ya "
                   "produce los datos para cerrar esta brecha; solo hay que leerlos juntos. Gracias. (5:10-5:55)")

ruta = OUT / "Pitch_El_miedo_que_no_baja_The_data_squad.pptx"
prs.save(ruta)
print("pitch:", ruta)
